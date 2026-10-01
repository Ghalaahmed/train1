"""
train_experiment.py
===================
Train / register models WITHOUT overwriting old ones. Every run gets its own folder:

    experiments/
    ├── experiments.xlsx              ← your tracking sheet (one row per experiment, filled automatically)
    ├── exp_001_original_v3/
    │   ├── model.pkl
    │   └── info.json                 ← settings + dataset + scores
    ├── exp_002_small_rf/
    └── ...
    current_model.txt                 ← which experiment the app uses (main_v3.py reads it)

Commands (run inside backend/, squat_env active):
    # 1) keep the model you already built as experiment #1 (it is copied, not changed)
    python train_experiment.py --register squat_model.pkl --name original_v3 --notes "first phases model"

    # 2) train a new experiment
    python train_experiment.py --name small_rf --trees 50 --depth 12 --notes "smaller + faster"
    python train_experiment.py --name tree8 --model tree --depth 8 --notes "single tree for the phone"

    # 3) choose which model the app uses (then restart uvicorn)
    python train_experiment.py --use exp_002_small_rf

    # 4) see all experiments
    python train_experiment.py --list

How every experiment is scored (same test for all → fair comparison):
  • split BY PERSON: 20% of people are only used for testing
  • macro F1 + balanced accuracy per joint (accuracy alone is misleading here)
  • rep accuracy: rep judged correct/incorrect like the camera (error >= 5 frames → rep fails)
  • speed: ms to predict one frame | size: MB
A new model is then trained on ALL the data and saved (scores come from the person split).
"""

import os, sys, json, time, shutil, pickle, hashlib, argparse, glob
import numpy as np
import pandas as pd
from sklearn.ensemble import RandomForestClassifier
from sklearn.tree import DecisionTreeClassifier
from sklearn.multioutput import MultiOutputClassifier
from sklearn.model_selection import GroupShuffleSplit
from sklearn.metrics import f1_score, balanced_accuracy_score

EXP_DIR   = "experiments"
SHEET     = os.path.join(EXP_DIR, "experiments.xlsx")
CURRENT   = "current_model.txt"
FEATURES  = ['knee_angle', 'hip_angle', 'back_angle', 'phase_id']
TARGETS   = ['knee_label', 'hip_label', 'back_label']
ERROR_FRAMES = 5

TRAIN_COLS = ["exp_id", "date", "name", "what changed", "dataset", "dataset hash", "rows", "people",
              "features", "model type", "trees", "max depth", "split",
              "F1 knee", "F1 hip", "F1 back", "F1 mean", "balanced acc mean", "rep accuracy",
              "ms / frame", "size MB", "folder", "decision (keep / discard)", "comments"]


# ───────────────────────── helpers ─────────────────────────
def file_hash(path):
    with open(path, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()[:8]


def next_id(name):
    os.makedirs(EXP_DIR, exist_ok=True)
    nums = [int(os.path.basename(p).split("_")[1]) for p in glob.glob(os.path.join(EXP_DIR, "exp_*"))
            if os.path.isdir(p) and os.path.basename(p).split("_")[1].isdigit()]
    return f"exp_{(max(nums) + 1 if nums else 1):03d}_{name}"


def single_thread(model):
    """n_jobs=-1 is fast for training but slow for predicting ONE frame → set 1 before saving."""
    for m in [model, *getattr(model, "estimators_", [])]:
        if hasattr(m, "n_jobs"):
            m.n_jobs = 1
    return model


def build_model(kind, trees, depth):
    if kind == "tree":
        return DecisionTreeClassifier(max_depth=depth, random_state=42)          # multi-output natively
    return MultiOutputClassifier(RandomForestClassifier(n_estimators=trees, max_depth=depth,
                                                        random_state=42, n_jobs=-1))


def rep_accuracy(df_test, pred):
    if 'rep_correct' not in df_test.columns:
        return None
    t = df_test.assign(err=(pred.max(axis=1) > 0))
    hits = []
    for _, g in t.sort_values('frame').groupby(['person_id', 'rep_id']):
        streak, ok = 0, 1
        for e in g.err.values:
            streak = streak + 1 if e else 0
            if streak >= ERROR_FRAMES:
                ok = 0; break
        hits.append(ok == g.rep_correct.iloc[0])
    return round(float(np.mean(hits)) * 100, 1)


def score(model, df, te):
    X, y = df.iloc[te][FEATURES], df.iloc[te][TARGETS]
    pred = np.asarray(model.predict(X))
    f1 = [f1_score(y.iloc[:, i], pred[:, i], average='macro', zero_division=0) * 100 for i in range(3)]
    bal = [balanced_accuracy_score(y.iloc[:, i], pred[:, i]) * 100 for i in range(3)]
    one = X.iloc[[0]]
    model.predict(one)
    t0 = time.perf_counter()
    for _ in range(50):
        model.predict(one)
    ms = (time.perf_counter() - t0) / 50 * 1000
    return dict(f1=[round(v, 1) for v in f1], f1_mean=round(np.mean(f1), 1),
                bal_mean=round(np.mean(bal), 1), rep_acc=rep_accuracy(df.iloc[te], pred), ms=round(ms, 1))


# ───────────────────────── tracking sheet ─────────────────────────
def make_sheet():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    head = Font(name="Arial", bold=True, color="FFFFFF")
    body = Font(name="Arial")
    fill_auto, fill_you = PatternFill("solid", fgColor="2F5D62"), PatternFill("solid", fgColor="B8860B")
    you = PatternFill("solid", fgColor="FFF2CC")

    # Training sheet — filled by this script
    ws = wb.active
    ws.title = "Training"
    ws.append(TRAIN_COLS)
    for i, c in enumerate(ws[1], 1):
        c.font, c.alignment = head, Alignment(wrap_text=True, vertical="center")
        c.fill = fill_you if TRAIN_COLS[i - 1].startswith(("decision", "comments")) else fill_auto
        ws.column_dimensions[c.column_letter].width = 14
    ws.column_dimensions["D"].width = 30; ws.column_dimensions["V"].width = 26; ws.column_dimensions["X"].width = 30
    ws.freeze_panes = "B2"
    dv = DataValidation(type="list", formula1='"keep,discard,in use"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("W2:W500")

    # Live tests sheet — you fill it after each camera test
    lt = wb.create_sheet("Live tests")
    cols = ["test_id", "date", "exp_id (model used)", "camera script", "lighting", "clothes",
            "distance / view", "scenario (what you did)", "reps planned", "system matched you",
            "match %", "false reps while setting up", "feedback stable? (Y/N)",
            "angles look right? (Y/N)", "session file", "what went wrong / notes"]
    lt.append(cols)
    for c in lt[1]:
        c.font, c.fill, c.alignment = head, fill_you, Alignment(wrap_text=True, vertical="center")
        lt.column_dimensions[c.column_letter].width = 15
    lt.column_dimensions["H"].width = 38; lt.column_dimensions["P"].width = 40
    example = ["T01 (example)", "2026-10-02", "exp_002_small_rf", "squat_camera_v4.py", "good", "fitted",
               "2.5 m, side", "5 correct, 3 shallow, 3 back lean", 11, 10, None, 0, "Y", "Y",
               "session_20261002_181500", "1 back-lean rep was counted as correct"]
    lt.append(example)
    for c in lt[2]:
        c.font = Font(name="Arial", italic=True, color="808080")
    for r in range(2, 301):
        lt[f"K{r}"] = f'=IFERROR(J{r}/I{r},"")'
        lt[f"K{r}"].number_format = "0%"
        for col in "ABCDEFGHIJLMNOP":
            if r > 2:
                lt[f"{col}{r}"].fill = you
                lt[f"{col}{r}"].font = body
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    lt.add_data_validation(yn); yn.add("M2:N300")
    lt.freeze_panes = "B2"

    # How to use
    hw = wb.create_sheet("How to use")
    for line in ["HOW TO USE THIS SHEET",
                 "",
                 "Training: one row per model. Filled AUTOMATICALLY by train_experiment.py — don't edit the dark-header columns.",
                 "  You fill only: 'decision (keep / discard)' and 'comments' (yellow header).",
                 "  All scores use the SAME person split, so rows can be compared fairly.",
                 "  F1 / balanced acc = frame level (does the model label each frame right?).",
                 "  rep accuracy = does the system judge the whole rep correct/incorrect like the dataset?",
                 "",
                 "Live tests: one row per camera test that YOU fill (yellow cells). Row 2 is an example.",
                 "  'match %' is calculated for you = system matched you / reps planned.",
                 "  Write which exp_id was in use (see current_model.txt) so results are linked to a model.",
                 "",
                 "Remember: the dataset is synthetic → high Training scores ≠ works on real people.",
                 "  The Live tests sheet is the real evidence."]:
        hw.append([line])
    hw.column_dimensions["A"].width = 110
    hw["A1"].font = Font(name="Arial", bold=True, size=13)
    for r in range(2, hw.max_row + 1):
        hw[f"A{r}"].font = body
    wb.save(SHEET)


def log_row(row):
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("openpyxl not installed → writing experiments.csv instead (pip install openpyxl for Excel)")
        path = os.path.join(EXP_DIR, "experiments.csv")
        pd.DataFrame([row], columns=TRAIN_COLS).to_csv(path, mode="a", header=not os.path.exists(path), index=False)
        return
    from openpyxl.styles import Font
    if not os.path.exists(SHEET):
        make_sheet()
    try:
        wb = load_workbook(SHEET)
        ws = wb["Training"]
        ws.append(row)
        for c in ws[ws.max_row]:
            c.font = Font(name="Arial")
        wb.save(SHEET)
    except PermissionError:
        print(f"!! {SHEET} is open in Excel — close it. Row saved to experiments.csv instead.")
        path = os.path.join(EXP_DIR, "experiments.csv")
        pd.DataFrame([row], columns=TRAIN_COLS).to_csv(path, mode="a", header=not os.path.exists(path), index=False)


# ───────────────────────── main actions ─────────────────────────
def run(args):
    df = pd.read_csv(args.dataset)
    tr, te = next(GroupShuffleSplit(n_splits=1, test_size=0.2, random_state=42).split(df, groups=df.person_id))
    exp_id = next_id(args.name)
    folder = os.path.join(EXP_DIR, exp_id)

    if args.register:                                   # existing model: copy + score only
        with open(args.register, "rb") as f:
            model = single_thread(pickle.load(f))
        kind = type(model).__name__
        trees = depth = ""
        try:
            base = model.estimators_[0]
            trees, depth = getattr(base, "n_estimators", ""), getattr(base, "max_depth", "")
        except Exception:
            pass
        s = score(model, df, te)
        split_note = "person 80/20 (model trained elsewhere — test people may have been seen)"
        os.makedirs(folder)
        shutil.copy(args.register, os.path.join(folder, "model.pkl"))
    else:                                               # new model
        print(f"Training {exp_id} ...")
        test_model = build_model(args.model, args.trees, args.depth).fit(df.iloc[tr][FEATURES], df.iloc[tr][TARGETS])
        s = score(single_thread(test_model), df, te)
        model = single_thread(build_model(args.model, args.trees, args.depth).fit(df[FEATURES], df[TARGETS]))
        kind = "RandomForest" if args.model == "rf" else "DecisionTree"
        trees, depth = (args.trees if args.model == "rf" else 1), args.depth
        split_note = "person 80/20 (seed 42)"
        os.makedirs(folder)
        with open(os.path.join(folder, "model.pkl"), "wb") as f:
            pickle.dump(model, f)

    size = round(os.path.getsize(os.path.join(folder, "model.pkl")) / 1e6, 2)
    info = dict(exp_id=exp_id, version=exp_id, date=time.strftime("%Y-%m-%d %H:%M"), name=args.name,
                notes=args.notes, dataset=args.dataset, dataset_hash=file_hash(args.dataset),
                rows=len(df), people=int(df.person_id.nunique()), features=FEATURES, targets=TARGETS,
                model_type=kind, trees=trees, max_depth=depth, split=split_note,
                scores=s, size_mb=size)
    with open(os.path.join(folder, "info.json"), "w", encoding="utf-8") as f:
        json.dump(info, f, indent=2, ensure_ascii=False, default=str)

    log_row([exp_id, info["date"], args.name, args.notes, args.dataset, info["dataset_hash"], len(df),
             info["people"], ", ".join(FEATURES), kind, trees, depth, split_note,
             *s["f1"], s["f1_mean"], s["bal_mean"], s["rep_acc"], s["ms"], size, folder, "", ""])

    print(f"\n{exp_id}")
    print(f"  F1 knee / hip / back : {s['f1'][0]} / {s['f1'][1]} / {s['f1'][2]}   (mean {s['f1_mean']})")
    print(f"  balanced acc (mean)  : {s['bal_mean']}")
    print(f"  rep accuracy         : {s['rep_acc']}")
    print(f"  speed / size         : {s['ms']} ms per frame | {size} MB")
    print(f"  saved                : {folder}/   + row in {SHEET}")
    print(f"\nTo use it in the app:  python train_experiment.py --use {exp_id}")


def use(exp_id):
    folder = os.path.join(EXP_DIR, exp_id)
    if not os.path.exists(os.path.join(folder, "model.pkl")):
        sys.exit(f"Not found: {folder}/model.pkl   (python train_experiment.py --list)")
    with open(CURRENT, "w") as f:
        f.write(folder)
    print(f"App now uses {exp_id}. Restart uvicorn (Ctrl+C, then run it again).")


def list_all():
    cur = open(CURRENT).read().strip() if os.path.exists(CURRENT) else None
    for p in sorted(glob.glob(os.path.join(EXP_DIR, "exp_*", "info.json"))):
        i = json.load(open(p, encoding="utf-8"))
        mark = "  <- in use" if cur and os.path.normpath(cur) == os.path.normpath(os.path.dirname(p)) else ""
        print(f"{i['exp_id']:28s} F1 {i['scores']['f1_mean']:5}  rep {i['scores']['rep_acc']}  "
              f"{i['scores']['ms']} ms  {i['size_mb']} MB  | {i['notes']}{mark}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="exp")
    ap.add_argument("--notes", default="")
    ap.add_argument("--dataset", default="squat_dataset_phases.csv")
    ap.add_argument("--model", choices=["rf", "tree"], default="rf")
    ap.add_argument("--trees", type=int, default=50)
    ap.add_argument("--depth", type=int, default=12)
    ap.add_argument("--register", help="existing .pkl to save as an experiment (copied, not retrained)")
    ap.add_argument("--use", help="exp_id the app should use")
    ap.add_argument("--list", action="store_true")
    a = ap.parse_args()
    if a.use:
        use(a.use)
    elif a.list:
        list_all()
    else:
        run(a)
